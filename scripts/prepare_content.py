"""Reorganize raw content/ files into the structure expected by convert.py.

Source layout (what you have):
  content/Content.docx               - master quest list
  content/{N}. {Title}.docx          - numbered quest docs (quests 1-N)

Target layout (what convert.py expects):
  content/quest-order.docx           - numbered list of all quest titles
  content/quests/{slug}.docx         - one doc per main quest
  content/side-quests/{slug}.docx    - one doc per side quest (if any)
"""
from __future__ import annotations
import re
import shutil
from pathlib import Path
from docx import Document


def slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_]+", "-", text)
    return text.strip("-")


def read_content_doc(path: Path):
    """Parse Content.docx -> (main_quests, side_quests) as lists of title strings."""
    doc = Document(str(path))
    main_quests: list[str] = []
    side_quests: list[str] = []
    in_side = False
    numbered = re.compile(r"^\d+[\.\)]\s+")
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        lower = text.lower()
        if "side quest" in lower and not numbered.match(text):
            in_side = True
            continue
        if not numbered.match(text):
            continue
        title = re.sub(r"^\d+[\.\)]\s*", "", text).strip()
        if not title:
            continue
        if in_side:
            side_quests.append(title)
        else:
            main_quests.append(title)
    return main_quests, side_quests


def build_quest_order_docx(main_quests: list[str], side_quests: list[str], out_path: Path) -> None:
    """Create a quest-order.docx with numbered paragraphs for all quests."""
    doc = Document()
    for i, title in enumerate(main_quests, start=1):
        doc.add_paragraph(f"{i}. {title}")
    if side_quests:
        doc.add_paragraph("Side Quests")
        for i, title in enumerate(side_quests, start=1):
            doc.add_paragraph(f"{i}. {title}")
    doc.save(str(out_path))
    print(f"  Created {out_path}")


def find_numbered_files(content_dir: Path) -> dict[int, Path]:
    """Map quest number -> Path for files named '{N}. {Title}.docx'."""
    mapping: dict[int, Path] = {}
    for p in content_dir.glob("*.docx"):
        m = re.match(r"^(\d+)\.", p.name)
        if m:
            mapping[int(m.group(1))] = p
    return mapping


def copy_quest_files(
    main_quests: list[str],
    numbered_files: dict[int, Path],
    quests_out: Path,
) -> None:
    quests_out.mkdir(parents=True, exist_ok=True)
    for i, title in enumerate(main_quests, start=1):
        slug = slugify(title)
        dest = quests_out / f"{slug}.docx"
        if i in numbered_files:
            shutil.copy2(numbered_files[i], dest)
            print(f"  [{i:2d}] {numbered_files[i].name!r} -> quests/{slug}.docx")
        else:
            print(f"  [{i:2d}] MISSING source for: {title!r}  (slug: {slug})")


def main() -> None:
    content_dir = Path("content")
    master_doc = content_dir / "Content.docx"

    if not master_doc.exists():
        raise FileNotFoundError(f"Expected {master_doc}")

    print("Reading Content.docx ...")
    main_quests, side_quests = read_content_doc(master_doc)
    print(f"  {len(main_quests)} main quests, {len(side_quests)} side quests")

    print("\nBuilding quest-order.docx ...")
    build_quest_order_docx(main_quests, side_quests, content_dir / "quest-order.docx")

    print("\nCopying numbered quest files -> content/quests/ ...")
    numbered = find_numbered_files(content_dir)
    copy_quest_files(main_quests, numbered, content_dir / "quests")

    print("\nDone. Run: python -m scripts.convert")


if __name__ == "__main__":
    main()
